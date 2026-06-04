"""Ingest Aarin_Complied_Stones.xlsx into stone_records.csv (+ parquet + QA report).

Run from the repo root:

    python src/01_ingest_excel.py [--limit N] [--validate-only] [--strict]
"""
from __future__ import annotations

import argparse
import json
import os
import sys
from collections import Counter
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import openpyxl
import pandas as pd


# ---------------------------------------------------------------------------
# Paths -- repo-relative only. The script must be run from the repo root.
# ---------------------------------------------------------------------------
SOURCE_XLSX    = Path("data/raw/Aarin_Complied_Stones.xlsx")
OUTPUT_CSV     = Path("data/processed/stone_records.csv")
OUTPUT_PARQUET = Path("data/processed/stone_records.parquet")
OUTPUT_REPORT  = Path("data/processed/data_quality_report.json")
DOTENV_PATH    = Path("config/.env")


# ---------------------------------------------------------------------------
# Workbook layout
# ---------------------------------------------------------------------------
SHEET_NAME = "Certificate Stock Report"
HEADER_ROW = 2
DATA_START = 3


# ---------------------------------------------------------------------------
# Grouping maps
# ---------------------------------------------------------------------------
SHAPE_GROUP = {
    "ROUND":         "round",
    "OVAL":          "oval",
    "OVAL STEP":     "oval",
    "HEART":         "heart",
    "EMERALD":       "emerald",
    "ASSCHER":       "asscher",
    "PEAR":          "pear",
    "PEAR STEP":     "pear",
    "MARQUISE":      "marquise",
    "PRINCESS":      "princess",
    "CUSHION BRLN":  "cushion",
    "CUSHION BRSQ":  "cushion",
    "CUSHION LN":    "cushion",
    "CUSHION SQ":    "cushion",
    "RADIANT BR":    "radiant",
}
# Everything else falls into "other" but shape_raw is preserved.

COLOR_GROUP = {
    "D": "colorless", "E": "colorless", "F": "colorless",
    "G": "near_colorless", "H": "near_colorless",
    "I": "slightly_tinted", "J": "slightly_tinted",
}

CLARITY_GROUP = {
    "FL":   "IF", "IF": "IF",
    "VVS1": "VVS", "VVS2": "VVS",
    "VS1":  "VS",  "VS2":  "VS",
    "SI1":  "SI",  "SI2":  "SI",
    "I1":   "I",
}

FLUOR_GROUP = {
    "NONE":        "none",
    "FAINT":       "faint",
    "MEDIUM":      "medium",
    "STRONG":      "strong",
    "VERY STRONG": "strong",
}

EYE_CLEAN_MAP = {"E-0": True, "E-A": False, "E-B": False, "N/A": None}

INCLUSION_VOCAB = {
    "feather":          "Feather",
    "crystal":          "Crystal",
    "needle":           "Needle",
    "pinpoint":         "Pinpoint",
    "cloud":            "Cloud",
    "indented natural": "Indented Natural",
    "twinning wisp":    "Twinning Wisp",
    "natural":          "Natural",
    "cavity":           "Cavity",
    "knot":             "Knot",
    "etch channel":     "Etch Channel",
    "extra facet":      "Extra Facet",
    "chip":             "Chip",
}

NULLISH = {None, "", "-", "n/a", "na"}

GIA_VERIFY_TPL = (
    "https://www.gia.edu/cs/Satellite?reportno={report_no}"
    "&childpagename=GIA/Page/ReportCheck"
    "&pagename=GIA/Dispatcher&c=Page&cid=1355954554547"
)


# ---------------------------------------------------------------------------
# Expected counts -- drive the final validation warnings
# ---------------------------------------------------------------------------
EXPECT_ROW_COUNT_APPROX   = 10537
EXPECT_MEDIA_COUNT_APPROX = 10498
EXPECT_PDF_COUNT_APPROX   = 10537
APPROX_TOL                = 5


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def _read_dotenv_key(path: Path, key: str) -> str | None:
    if not path.exists():
        return None
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        k, _, v = line.partition("=")
        if k.strip() == key:
            v = v.strip().strip('"').strip("'")
            return v or None
    return None


def clean(v):
    """Map blanks / '-' / 'N/A' / whitespace-only to None; pass everything else through."""
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        if s.lower() in NULLISH:
            return None
        return s
    return v


def to_float(v):
    v = clean(v)
    if v is None:
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def parse_inclusions(comment: str | None):
    """Return (recognized_set, unknown_list) from the Comment cell."""
    if not comment:
        return set(), []
    recognized: set[str] = set()
    unknown: list[str] = []
    for raw_tok in str(comment).split(","):
        tok = raw_tok.strip()
        if not tok:
            continue
        canonical = INCLUSION_VOCAB.get(tok.lower())
        if canonical is not None:
            recognized.add(canonical)
        else:
            unknown.append(tok)
    return recognized, unknown


def weak_label(milky, blkc, blks, tabinc) -> str:
    vals = [v for v in (milky, blkc, blks, tabinc) if v]
    if not vals:
        return "unknown"
    # Codes end with the severity character (e.g. M-0, N-TA, TB-C). Take the last char.
    severities = [str(v).rstrip().upper()[-1] for v in vals]
    if any(s in {"B", "C", "D"} for s in severities):
        return "visible"
    if any(s == "A" for s in severities):
        return "slight"
    if all(s == "0" for s in severities):
        return "clean"
    return "unknown"


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def build_header_index(ws) -> dict[str, int]:
    headers: dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        name = ws.cell(row=HEADER_ROW, column=c).value
        if name:
            headers[str(name).strip()] = c
    return headers


REQUIRED_HEADERS = [
    "SrNo", "RefNo", "Shape", "Weight", "Color", "Purity", "Cut", "Polish", "Symn",
    "Fluor", "RRate", "GlobalNo", "Mes1", "Mes2", "Mes3", "Ratio", "DepthPer", "Table",
    "C/H", "CrAng", "P/D", "IsBGM", "Cert", "MediaLink", "PdfLink", "HnALink",
    "EYEClean", "Comment", "ReportCmnt", "CertificateNo", "ReportNo", "Inscription",
    "ColInt", "BlkC", "BlkS", "TableInclusion", "Milky", "Remark", "Girdle",
    "StarLength", "LowerHalf", "CuletSize", "GMin", "GMax", "Days", "DiamType",
]


def ingest(args) -> tuple[list[dict], dict]:
    if not SOURCE_XLSX.exists():
        sys.exit(f"ERROR: source xlsx not found at {SOURCE_XLSX} (run from repo root)")

    wb = openpyxl.load_workbook(SOURCE_XLSX, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        sys.exit(f"ERROR: sheet '{SHEET_NAME}' not in workbook; found {wb.sheetnames}")
    ws = wb[SHEET_NAME]

    H = build_header_index(ws)
    missing = [h for h in REQUIRED_HEADERS if h not in H]
    if missing:
        sys.exit(f"ERROR: missing expected headers in row {HEADER_ROW}: {missing}")

    def cell(r: int, name: str):
        return ws.cell(row=r, column=H[name])

    def val(r: int, name: str):
        return clean(cell(r, name).value)

    def hyperlink_target(r: int, name: str):
        link = cell(r, name).hyperlink
        return link.target if link is not None else None

    last_row = ws.max_row
    if args.limit is not None:
        last_row = min(last_row, DATA_START + args.limit - 1)

    rows: list[dict] = []
    shape_group_counts:    Counter = Counter()
    color_group_counts:    Counter = Counter()
    clarity_group_counts:  Counter = Counter()
    fluor_group_counts:    Counter = Counter()
    inclusion_type_counts: Counter = Counter()
    raw_shape_other_examples: Counter = Counter()
    unknown_inclusion_examples: Counter = Counter()
    warnings_depth_under_40: list[dict] = []
    warnings_missing_required: list[dict] = []

    for r in range(DATA_START, last_row + 1):
        stone_id = val(r, "RefNo")
        if not stone_id:
            continue  # defensive; no blanks observed

        # Raw labels
        shape_raw           = val(r, "Shape")
        color_raw           = val(r, "Color")
        clarity_raw         = val(r, "Purity")
        cut_raw             = val(r, "Cut")
        polish_raw          = val(r, "Polish")
        symn_raw            = val(r, "Symn")
        fluor_raw           = val(r, "Fluor")
        eye_clean_raw       = val(r, "EYEClean")
        milky_raw           = val(r, "Milky")
        blk_c_raw           = val(r, "BlkC")
        blk_s_raw           = val(r, "BlkS")
        table_inclusion_raw = val(r, "TableInclusion")
        culet_raw           = val(r, "CuletSize")
        girdle_min_raw      = val(r, "GMin")
        girdle_max_raw      = val(r, "GMax")
        col_int_raw         = val(r, "ColInt")
        is_bgm_raw          = val(r, "IsBGM")
        inscription_raw     = val(r, "Inscription")
        diam_type_raw       = val(r, "DiamType")

        # Free text -- keep verbatim, including any internal whitespace
        comment        = clean(cell(r, "Comment").value)
        report_comment = clean(cell(r, "ReportCmnt").value)
        remark         = clean(cell(r, "Remark").value)

        # Links
        media_link = hyperlink_target(r, "MediaLink")
        pdf_link   = hyperlink_target(r, "PdfLink")
        hna_link   = hyperlink_target(r, "HnALink")
        media_ref: str | None = None
        if media_link:
            try:
                media_ref = parse_qs(urlparse(media_link).query).get("r", [None])[0]
            except Exception:
                media_ref = None

        # Identity
        certificate_no = val(r, "CertificateNo")
        report_no      = val(r, "ReportNo")
        gia_verify_url = (
            GIA_VERIFY_TPL.format(report_no=report_no) if report_no else None
        )

        # Measurements
        weight_ct        = to_float(cell(r, "Weight").value)
        mes_length       = to_float(cell(r, "Mes1").value)
        mes_width        = to_float(cell(r, "Mes2").value)
        mes_depth        = to_float(cell(r, "Mes3").value)
        ratio            = to_float(cell(r, "Ratio").value)
        depth_pct        = to_float(cell(r, "DepthPer").value)
        table_pct        = to_float(cell(r, "Table").value)
        crown_angle      = to_float(cell(r, "CrAng").value)
        pavilion_depth   = to_float(cell(r, "P/D").value)
        star_length_pct  = to_float(cell(r, "StarLength").value)
        lower_half_pct   = to_float(cell(r, "LowerHalf").value)

        # Commercial
        rap_rate      = to_float(cell(r, "RRate").value)
        rap_total     = (weight_ct * rap_rate) if (weight_ct is not None and rap_rate is not None) else None
        days_in_stock = to_float(cell(r, "Days").value)

        # Grouped labels (raw is the source of truth; group falls through to defaults)
        shape_group   = SHAPE_GROUP.get(shape_raw.upper() if shape_raw else "", "other")
        if shape_group == "other" and shape_raw:
            raw_shape_other_examples[shape_raw] += 1

        color_group   = COLOR_GROUP.get(color_raw.upper() if color_raw else "", "tinted_or_other")
        clarity_group = CLARITY_GROUP.get(clarity_raw.upper() if clarity_raw else "")
        fluor_group   = FLUOR_GROUP.get(fluor_raw.upper() if fluor_raw else "")

        eye_clean = EYE_CLEAN_MAP.get(eye_clean_raw) if eye_clean_raw else None

        # Inclusions
        recognized, unknown = parse_inclusions(comment)
        inclusion_types_pipe = "|".join(sorted(recognized))
        inclusion_types_json = json.dumps(sorted(recognized))
        unknown_inclusion_terms = "|".join(sorted({u for u in unknown}))
        for tok in recognized:
            inclusion_type_counts[tok] += 1
        for tok in unknown:
            unknown_inclusion_examples[tok] += 1

        # Inclusion weak label
        inclusion_weak_label = weak_label(milky_raw, blk_c_raw, blk_s_raw, table_inclusion_raw)

        # Counters
        shape_group_counts[shape_group] += 1
        color_group_counts[color_group] += 1
        if clarity_group:
            clarity_group_counts[clarity_group] += 1
        if fluor_group:
            fluor_group_counts[fluor_group] += 1

        # Warning lists
        if depth_pct is not None and depth_pct < 40:
            warnings_depth_under_40.append(
                {"excel_row_num": r, "stone_id": stone_id, "depth_pct": depth_pct}
            )
        missing_required = [
            name for name, v in (
                ("stone_id", stone_id),
                ("shape_raw", shape_raw),
                ("color_raw", color_raw),
                ("clarity_raw", clarity_raw),
            )
            if v in (None, "")
        ]
        if missing_required:
            warnings_missing_required.append(
                {"excel_row_num": r, "stone_id": stone_id, "missing": missing_required}
            )

        rows.append({
            # Identity & provenance
            "excel_row_num":   r,
            "stone_id":        stone_id,
            "global_no":       val(r, "GlobalNo"),
            "cert_lab":        val(r, "Cert"),
            "certificate_no":  certificate_no,
            "report_no":       report_no,

            # Links
            "media_link":      media_link,
            "media_ref":       media_ref,
            "pdf_link":        pdf_link,
            "hna_link":        hna_link,
            "gia_verify_url":  gia_verify_url,
            "has_media":       media_link is not None,
            "has_pdf":         pdf_link is not None,

            # Raw labels (preserved verbatim)
            "shape_raw":              shape_raw,
            "color_raw":              color_raw,
            "clarity_raw":            clarity_raw,
            "cut_raw":                cut_raw,
            "polish_raw":             polish_raw,
            "symn_raw":               symn_raw,
            "fluor_raw":              fluor_raw,
            "eye_clean_raw":          eye_clean_raw,
            "milky_raw":              milky_raw,
            "blk_c_raw":              blk_c_raw,
            "blk_s_raw":              blk_s_raw,
            "table_inclusion_raw":    table_inclusion_raw,
            "culet_raw":              culet_raw,
            "girdle_min_raw":         girdle_min_raw,
            "girdle_max_raw":         girdle_max_raw,
            "col_int_raw":            col_int_raw,
            "is_bgm_raw":             is_bgm_raw,
            "inscription_raw":        inscription_raw,
            "diam_type_raw":          diam_type_raw,
            "comment":                comment,
            "report_comment":         report_comment,
            "remark":                 remark,

            # Grouped labels (model-ready)
            "shape_group":            shape_group,
            "color_group":            color_group,
            "clarity_group":          clarity_group,
            "fluorescence_group":     fluor_group,
            "eye_clean":              eye_clean,
            "inclusion_weak_label":   inclusion_weak_label,

            # Inclusion fields
            "inclusion_types":         inclusion_types_pipe,
            "inclusion_types_json":    inclusion_types_json,
            "unknown_inclusion_terms": unknown_inclusion_terms,

            # Measurements
            "weight_ct":       weight_ct,
            "mes_length":      mes_length,
            "mes_width":       mes_width,
            "mes_depth":       mes_depth,
            "ratio":           ratio,
            "depth_pct":       depth_pct,
            "table_pct":       table_pct,
            "crown_angle":     crown_angle,
            "pavilion_depth":  pavilion_depth,
            "star_length_pct": star_length_pct,
            "lower_half_pct":  lower_half_pct,

            # Commercial
            "rap_rate":        rap_rate,
            "rap_total":       rap_total,
            "days_in_stock":   days_in_stock,
        })

    report = {
        "row_count":                 len(rows),
        "unique_stone_id_count":     len({row["stone_id"] for row in rows}),
        "has_media_count":           sum(1 for row in rows if row["has_media"]),
        "has_pdf_count":             sum(1 for row in rows if row["has_pdf"]),
        "missing_media_count":       sum(1 for row in rows if not row["has_media"]),
        "shape_group_counts":        dict(shape_group_counts.most_common()),
        "color_group_counts":        dict(color_group_counts.most_common()),
        "clarity_group_counts":      dict(clarity_group_counts.most_common()),
        "fluorescence_group_counts": dict(fluor_group_counts.most_common()),
        "inclusion_type_counts":     dict(inclusion_type_counts.most_common()),
        "raw_shape_other_examples":  dict(raw_shape_other_examples.most_common()),
        "unknown_inclusion_examples": dict(unknown_inclusion_examples.most_common()),
        "warning_rows_depth_under_40":         warnings_depth_under_40,
        "warning_rows_missing_required_fields": warnings_missing_required,
    }
    return rows, report


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------
def run_validations(report: dict, args) -> list[str]:
    warns: list[str] = []
    n = report["row_count"]

    if args.limit is None:
        if abs(n - EXPECT_ROW_COUNT_APPROX) > APPROX_TOL:
            warns.append(f"row_count={n} differs from expected ~{EXPECT_ROW_COUNT_APPROX} by more than ±{APPROX_TOL}")
        if abs(report["has_media_count"] - EXPECT_MEDIA_COUNT_APPROX) > APPROX_TOL:
            warns.append(
                f"has_media_count={report['has_media_count']} differs from expected ~{EXPECT_MEDIA_COUNT_APPROX}"
            )
        if abs(report["has_pdf_count"] - EXPECT_PDF_COUNT_APPROX) > APPROX_TOL:
            warns.append(
                f"has_pdf_count={report['has_pdf_count']} differs from expected ~{EXPECT_PDF_COUNT_APPROX}"
            )

    if report["unique_stone_id_count"] != n:
        warns.append(
            f"stone_id not unique: {report['unique_stone_id_count']} unique vs {n} rows"
        )

    if report["warning_rows_missing_required_fields"]:
        warns.append(
            f"{len(report['warning_rows_missing_required_fields'])} rows missing a required field "
            f"(see warning_rows_missing_required_fields)"
        )

    if report["unknown_inclusion_examples"]:
        sample = list(report["unknown_inclusion_examples"].items())[:5]
        warns.append(f"unknown inclusion tokens encountered: {sample}")

    if report["raw_shape_other_examples"]:
        # Informational, not surfaced as WARN unless --strict and you want to gate on it.
        # We keep it INFO so the default run is clean.
        pass

    return warns


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def parse_args():
    p = argparse.ArgumentParser(description="Ingest Aarin xlsx -> stone_records.csv")
    p.add_argument("--limit", type=int, default=None, help="Only read the first N data rows")
    p.add_argument("--validate-only", action="store_true", help="Read+validate, write nothing")
    p.add_argument("--strict", action="store_true", help="Exit nonzero on any validation warning")
    return p.parse_args()


def main():
    args = parse_args()

    media_token = (
        os.environ.get("AARIN_MEDIA_TOKEN")
        or _read_dotenv_key(DOTENV_PATH, "AARIN_MEDIA_TOKEN")
    )
    print(f"INFO: input = {SOURCE_XLSX}")
    print(f"INFO: outputs = {OUTPUT_CSV}, {OUTPUT_PARQUET} (optional), {OUTPUT_REPORT}")
    print(f"INFO: AARIN_MEDIA_TOKEN: {'present' if media_token else 'absent'}")
    if args.limit:
        print(f"INFO: --limit {args.limit}")
    if args.validate_only:
        print("INFO: --validate-only (no files will be written)")

    rows, report = ingest(args)

    # Headline
    print(
        f"\nREAD {report['row_count']} rows | "
        f"media={report['has_media_count']} | pdf={report['has_pdf_count']} | "
        f"unique_stone_id={report['unique_stone_id_count']}"
    )
    print(f"shape_groups: {report['shape_group_counts']}")
    print(f"color_groups: {report['color_group_counts']}")
    print(f"clarity_groups: {report['clarity_group_counts']}")
    print(f"fluor_groups:   {report['fluorescence_group_counts']}")

    # Validations
    warns = run_validations(report, args)
    for w in warns:
        print(f"WARN: {w}")

    if report["raw_shape_other_examples"]:
        print(f"INFO: shape_raw values that fell into 'other' group: {report['raw_shape_other_examples']}")
    if report["warning_rows_depth_under_40"]:
        print(f"INFO: {len(report['warning_rows_depth_under_40'])} rows with depth_pct<40 (kept; flagged in QA report)")

    # Writes
    if not args.validate_only:
        OUTPUT_CSV.parent.mkdir(parents=True, exist_ok=True)
        df = pd.DataFrame(rows)
        df.to_csv(OUTPUT_CSV, index=False, encoding="utf-8")
        print(f"WROTE {OUTPUT_CSV} ({df.shape[0]} rows x {df.shape[1]} cols)")

        try:
            import pyarrow  # noqa: F401
            df.to_parquet(OUTPUT_PARQUET, index=False)
            print(f"WROTE {OUTPUT_PARQUET}")
        except ImportError:
            print("INFO: pyarrow not installed; skipping parquet output")

        OUTPUT_REPORT.write_text(json.dumps(report, indent=2, default=str), encoding="utf-8")
        print(f"WROTE {OUTPUT_REPORT}")

    if args.strict and warns:
        print(f"STRICT: {len(warns)} warning(s); exiting with status 1")
        sys.exit(1)


if __name__ == "__main__":
    main()
